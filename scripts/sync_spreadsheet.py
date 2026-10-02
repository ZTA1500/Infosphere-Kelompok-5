"""Write the "Ruangan (peta)" sheet of data/qa/Dataset Jarvis.xlsx from data/rooms.json.

    python scripts/sync_spreadsheet.py

The sheet lists every room on the floor maps with its answer in Indonesian,
English and Mandarin — exactly what the chatbot says. It is a reference for
the team: the chatbot reads room answers from data/rooms.json (edit the room
there, then run this script to refresh the sheet). Only the FIRST sheet of
the workbook is used for training, so this sheet never changes the model.
"""
import os
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from server.dataset import load_rooms  # noqa: E402
from server.i18n import table  # noqa: E402
from server.paths import MAIN_DATASET  # noqa: E402
from server.qa_locations import room_entries  # noqa: E402

SHEET = 'Ruangan (peta)'
COLUMNS = [('ID', 16), ('Lantai', 8), ('Nama', 24), ('Name (English)', 26), ('名称 (中文)', 22),
           ('Jawaban (Indonesia)', 48), ('Answer (English)', 48), ('回答 (中文)', 36), ('Contoh pertanyaan', 30)]


def main():
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass
    rooms = load_rooms()['rooms']
    translations = table()
    wb = load_workbook(MAIN_DATASET)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)          # always after the training sheet

    ws['A1'] = ('Dibuat otomatis dari data/rooms.json — jangan diedit di sini. Ubah ruangan di rooms.json, '
                'lalu jalankan: python scripts/sync_spreadsheet.py  (Generated from rooms.json; this sheet is '
                'not used for training.)')
    ws['A1'].font = Font(italic=True, color='555555')
    head = 3
    for col, (title, width) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=head, column=col, value=title)
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='1565C0')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width

    for _i, room, display, answer in room_entries(rooms):
        tr = translations.get(answer.strip(), {})
        ws.append([room['id'], room['floor'], display, room.get('name_en') or room['name'],
                   room.get('name_zh') or room['name'], answer, tr.get('en', ''), tr.get('zh', ''),
                   f'Dimana lokasi {display}? / Where is {room.get("name_en") or display}? / '
                   f'{room.get("name_zh") or display}在哪里？'])
        for c in ws[ws.max_row]:
            c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = ws.cell(row=head + 1, column=1)
    ws.auto_filter.ref = f'A{head}:{ws.cell(row=ws.max_row, column=len(COLUMNS)).coordinate}'
    wb.save(MAIN_DATASET)
    print(f'Wrote {len(rooms)} rooms to the "{SHEET}" sheet of {os.path.basename(MAIN_DATASET)}')


if __name__ == '__main__':
    main()
